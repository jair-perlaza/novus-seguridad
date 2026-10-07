"""
Generate professional PDF/HTML/JSON/CSV exports for NOVUS security reports.
"""
import csv
import hashlib
import io
import json
import os
import re

from utils.logger import logger

REPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "reports")
NOVUS_SLOGAN = "Protección inteligente en tiempo real"
NOVUS_BRAND = "NOVUS Security Platform"


from utils.pdf_text import normalize_pdf_text


def _safe_text(value):
    if value is None:
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ")
    return normalize_pdf_text(text)


def _write_section(pdf, title, content):
    if not content:
        return
    pdf.set_x(pdf.l_margin)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(pdf.epw, 7, _safe_text(title), ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(pdf.epw, 6, _safe_text(content))
    pdf.ln(2)


def _write_list(pdf, title, items):
    if not items:
        return
    pdf.set_x(pdf.l_margin)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(pdf.epw, 7, _safe_text(title), ln=True)
    pdf.set_font("Helvetica", "", 10)
    for idx, item in enumerate(items, 1):
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(pdf.epw, 6, _safe_text(f"{idx}. {item}"))
    pdf.ln(2)


def _footer(pdf, page_num, total_hint=""):
    pdf.set_y(-15)
    pdf.set_font("Helvetica", "I", 8)
    pdf.set_text_color(120, 120, 120)
    pdf.cell(0, 5, _safe_text(f"{NOVUS_SLOGAN}  |  Pagina {page_num}"), align="C")
    pdf.set_text_color(0, 0, 0)


def _draw_bar_chart(pdf, title, data: dict, max_width=170):
    """Gráfico de barras horizontal con datos reales."""
    if not data:
        return
    pdf.ln(3)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(pdf.epw, 7, _safe_text(title), ln=True)
    pdf.set_font("Helvetica", "", 9)
    numeric = {k: v for k, v in data.items() if isinstance(v, (int, float))}
    if not numeric:
        _write_section(pdf, "", "Sin datos disponibles para grafico.")
        return
    max_val = max(numeric.values()) or 1
    colors = [(239, 68, 68), (249, 115, 22), (234, 179, 8), (34, 197, 94), (148, 163, 184)]
    for i, (label, val) in enumerate(numeric.items()):
        bar_w = max(2, (val / max_val) * max_width)
        pdf.set_x(pdf.l_margin)
        pdf.cell(35, 6, _safe_text(label[:12]), ln=0)
        r, g, b = colors[i % len(colors)]
        pdf.set_fill_color(r, g, b)
        pdf.cell(bar_w, 6, "", fill=True, ln=0)
        pdf.cell(10, 6, str(val), ln=True)
    pdf.ln(4)


def generate_enterprise_pdf(report: dict) -> str:
    """PDF empresarial con portada, gráficos y firma."""
    try:
        from fpdf import FPDF  # noqa: F401 — comprobación de dependencia
    except ImportError:
        return _generate_enterprise_html(report)

    from utils.novus_fpdf import create_novus_pdf

    pdf = create_novus_pdf()
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(auto=True, margin=20)
    meta = report.get("metadata") or {}
    cover = report.get("cover") or {}
    exec_sum = report.get("executive_summary") or {}
    charts = report.get("charts") or {}
    sig = report.get("signature") or {}

    # --- PORTADA ---
    pdf.add_page()
    pdf.set_fill_color(93, 95, 239)
    pdf.rect(0, 0, 210, 55, "F")
    pdf.set_text_color(255, 255, 255)
    pdf.set_y(18)
    pdf.set_font("Helvetica", "B", 28)
    pdf.cell(pdf.epw, 12, "NOVUS", ln=True, align="C")
    pdf.set_font("Helvetica", "", 11)
    pdf.cell(pdf.epw, 7, _safe_text(NOVUS_BRAND), ln=True, align="C")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(20)
    pdf.set_font("Helvetica", "B", 20)
    pdf.cell(pdf.epw, 10, _safe_text(cover.get("titulo", "Informe de Seguridad")), ln=True, align="C")
    pdf.ln(8)
    pdf.set_font("Helvetica", "", 11)
    for label, val in [
        ("Empresa", cover.get("empresa")),
        ("Sector", cover.get("sector")),
        ("Fecha", cover.get("fecha")),
        ("Hora", cover.get("hora")),
        ("Informe N.", report.get("id")),
        ("Version", report.get("version")),
        ("Clasificacion", report.get("classification")),
    ]:
        pdf.cell(pdf.epw, 7, _safe_text(f"{label}: {val}"), ln=True, align="C")
    pdf.ln(10)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(pdf.epw, 8, f"Nivel de seguridad: {_safe_text(cover.get('nivel_seguridad'))}", ln=True, align="C")
    pdf.ln(6)
    pdf.set_font("Helvetica", "I", 10)
    pdf.multi_cell(pdf.epw, 6, _safe_text(cover.get("resumen_ejecutivo_breve", "")), align="C")
    _footer(pdf, 1)

    # --- RESUMEN EJECUTIVO ---
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(pdf.epw, 10, "Resumen Ejecutivo", ln=True)
    pdf.ln(4)
    for label, key in [
        ("Estado general", "estado_general"),
        ("Problema mas importante", "problema_mas_importante"),
        ("Vulnerabilidades", "vulnerabilidades"),
        ("Incidentes", "incidentes"),
        ("Amenazas", "amenazas"),
        ("Dispositivos monitoreados", "dispositivos_monitoreados"),
        ("Recomendacion principal", "recomendacion_principal"),
    ]:
        _write_section(pdf, label, str(exec_sum.get(key, "")))
    gs = report.get("general_status") or {}
    _write_section(pdf, "Nivel de riesgo", gs.get("nivel_riesgo"))
    _write_section(pdf, "Nivel de confianza", gs.get("nivel_confianza"))
    _footer(pdf, 2)

    # --- GRAFICOS ---
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(pdf.epw, 10, "Indicadores (datos reales)", ln=True)
    _draw_bar_chart(pdf, "Distribucion de riesgos", charts.get("distribucion_riesgos") or {})
    _draw_bar_chart(pdf, "Estado de infraestructura", charts.get("estado_infraestructura") or {})
    _draw_bar_chart(pdf, "Evolucion vulnerabilidades", charts.get("evolucion_vulnerabilidades") or {})
    _draw_bar_chart(pdf, "Uso de recursos (%)", charts.get("uso_recursos") or {})
    _footer(pdf, 3)

    # --- DETALLE TECNICO ---
    details = report.get("technical_details") or []
    if details:
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(pdf.epw, 10, "Detalle Tecnico de Hallazgos", ln=True)
        for idx, d in enumerate(details[:20], 1):
            pdf.ln(3)
            pdf.set_font("Helvetica", "B", 12)
            pdf.cell(pdf.epw, 8, _safe_text(f"{idx}. {d.get('id', 'Hallazgo')}"), ln=True)
            pdf.set_font("Helvetica", "", 9)
            for lbl, key in [
                ("Descripcion", "descripcion"), ("Motor", "motor"), ("Confianza", "confianza"),
                ("Riesgo", "riesgo"), ("Estado", "estado"), ("Impacto", "impacto"),
                ("Fecha", "fecha"), ("Hora", "hora"),
            ]:
                pdf.set_x(pdf.l_margin)
                pdf.multi_cell(pdf.epw, 5, _safe_text(f"{lbl}: {d.get(key, '')}"))
            ev = d.get("evidencia") or []
            if ev:
                pdf.set_x(pdf.l_margin)
                pdf.multi_cell(pdf.epw, 5, _safe_text(f"Evidencia: {'; '.join(str(e) for e in ev[:3])}"))
            recs = d.get("recomendaciones") or []
            if recs:
                pdf.set_x(pdf.l_margin)
                pdf.multi_cell(pdf.epw, 5, _safe_text(f"Recomendaciones: {' | '.join(recs[:2])}"))
            if pdf.get_y() > 250:
                _footer(pdf, pdf.page_no())
                pdf.add_page()

    # --- HISTORIAL ---
    hist = report.get("history_evolution") or {}
    if hist:
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(pdf.epw, 10, "Historial y Evolucion", ln=True)
        _write_section(pdf, "Informes anteriores", str(hist.get("informes_anteriores", 0)))
        _write_section(pdf, "Vulnerabilidades corregidas", str(hist.get("vulnerabilidades_corregidas", 0)))
        _write_section(pdf, "Tendencia de riesgo", hist.get("riesgo_tendencia"))
        _write_list(pdf, "Mejoras", hist.get("mejoras") or [])

    # --- FIRMA ---
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(pdf.epw, 10, "Firma del Informe", ln=True)
    pdf.ln(6)
    pdf.set_font("Helvetica", "", 11)
    for lbl, key in [
        ("Informe generado por", "generado_por"),
        ("Fecha", "fecha"),
        ("Hora", "hora"),
        ("Version", "version"),
        ("Operador", "operador"),
        ("Integridad SHA-256", "integridad_sha256"),
    ]:
        pdf.cell(pdf.epw, 8, _safe_text(f"{lbl}: {sig.get(key, '')}"), ln=True)
    pdf.ln(10)
    pdf.set_font("Helvetica", "I", 10)
    pdf.multi_cell(pdf.epw, 6, _safe_text(NOVUS_SLOGAN))
    _footer(pdf, pdf.page_no())

    out_path = os.path.join(REPORTS_DIR, f"{report['id']}.pdf")
    pdf.safe_output(out_path)
    return out_path


def _generate_manual_defense_pdf(report):
    """PDF para informes MDR del Centro de Defensa."""
    try:
        from fpdf import FPDF
    except ImportError:
        from services.manual_defense_report_html import render_manual_defense_report_html

        out_path = os.path.join(REPORTS_DIR, f"{report['id']}.html")
        with open(out_path, "w", encoding="utf-8") as handle:
            handle.write(render_manual_defense_report_html(report))
        logger.warning("fpdf2 no instalado — informe MDR exportado como HTML")
        return out_path

    tech = report.get("technical") or {}
    from services.manual_defense_report_html import ensure_defense_report

    dr = ensure_defense_report(report)
    from utils.novus_fpdf import create_novus_pdf

    pdf = create_novus_pdf()
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(pdf.epw, 10, _safe_text("NOVUS - Centro de Defensa"), ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(pdf.epw, 6, _safe_text(f"ID: {report.get('id')} | {report.get('fecha')}"), ln=True)
    pdf.cell(
        pdf.epw, 6,
        _safe_text(f"{dr.get('mechanism_title', '')} | Severidad: {report.get('severidad')}"),
        ln=True,
    )
    pdf.ln(4)
    _write_section(pdf, "Resumen ejecutivo", dr.get("executive_summary") or report.get("resumen_ejecutivo"))
    for sec in dr.get("sections") or []:
        content = sec.get("content")
        if isinstance(content, (dict, list)):
            content = normalize_pdf_text(json.dumps(content, ensure_ascii=False)[:4000])
        _write_section(pdf, sec.get("title") or sec.get("key"), content)
    nf = dr.get("not_found_verified") or []
    if nf:
        _write_list(pdf, "No se encontro (amenazas verificadas)", [x.get("label") for x in nf])
    kernel = dr.get("kernel_analysis") or {}
    _write_section(pdf, "Kernel IA - resumen", kernel.get("summary"))
    if kernel.get("risks_found"):
        lines = [
            f"{r.get('label')} ({r.get('risk')})"
            for r in kernel.get("risks_found")[:20]
            if isinstance(r, dict)
        ]
        _write_list(pdf, "Riesgos encontrados", lines)
    if kernel.get("risks_not_found_verified"):
        _write_list(
            pdf,
            "Riesgos no encontrados (verificados)",
            [x.get("label") for x in kernel.get("risks_not_found_verified") if isinstance(x, dict)],
        )
    _write_list(pdf, "Recomendaciones", kernel.get("recommendations") or [])
    _write_list(pdf, "Prioridades", kernel.get("priorities") or [])
    _write_section(pdf, "Conclusion", dr.get("conclusion") or report.get("conclusiones"))
    out_path = os.path.join(REPORTS_DIR, f"{report['id']}.pdf")
    pdf.safe_output(out_path)
    return out_path


def _write_table_header(pdf, headers, col_widths):
    pdf.set_font("Helvetica", "B", 9)
    pdf.set_fill_color(241, 245, 249)
    x0 = pdf.l_margin
    y0 = pdf.get_y()
    for i, (h, w) in enumerate(zip(headers, col_widths)):
        pdf.set_xy(x0 + sum(col_widths[:i]), y0)
        pdf.cell(w, 7, _safe_text(h), border=1, fill=True)
    pdf.ln(7)


def _write_table_row(pdf, cells, col_widths):
    pdf.set_font("Helvetica", "", 8)
    x0 = pdf.l_margin
    y0 = pdf.get_y()
    for i, (c, w) in enumerate(zip(cells, col_widths)):
        pdf.set_xy(x0 + sum(col_widths[:i]), y0)
        pdf.cell(w, 6, _safe_text(str(c)[:48]), border=1)
    pdf.ln(6)


def generate_presentation_pdf(report: dict) -> str:
    """PDF alineado a la vista ejecutiva (sin JSON ni campos internos)."""
    from services.report_presentation_service import build_executive_presentation

    pres = build_executive_presentation(report)
    try:
        from fpdf import FPDF  # noqa: F401
    except ImportError:
        return _generate_presentation_html(pres, report)

    from utils.novus_fpdf import create_novus_pdf

    pdf = create_novus_pdf()
    pdf.set_margins(14, 14, 14)
    pdf.set_auto_page_break(auto=True, margin=18)
    exec_b = pres.get("executive") or {}
    scan = pres.get("scan_summary") or {}

    pdf.add_page()
    pdf.set_fill_color(93, 95, 239)
    pdf.rect(0, 0, 210, 42, "F")
    pdf.set_text_color(255, 255, 255)
    pdf.set_y(12)
    pdf.set_font("Helvetica", "B", 22)
    pdf.cell(pdf.epw, 10, "NOVUS", ln=True, align="C")
    pdf.set_font("Helvetica", "", 11)
    pdf.cell(pdf.epw, 7, _safe_text(pres.get("title") or "Informe de Seguridad"), ln=True, align="C")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(14)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(pdf.epw, 8, _safe_text(f"Informe {report.get('id')}"), ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(pdf.epw, 6, _safe_text(f"Tipo: {pres.get('tipo')}"), ln=True)
    _footer(pdf, 1)

    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(pdf.epw, 9, "Resumen ejecutivo", ln=True)
    pdf.ln(2)
    for label, key in (
        ("Fecha", "fecha"), ("Hora", "hora"), ("Usuario", "usuario"), ("Equipo", "equipo"),
        ("Red analizada", "red_analizada"), ("Estado general", "estado_general"),
        ("Nivel de riesgo", "nivel_riesgo"), ("Prioridad", "prioridad"),
    ):
        pdf.set_font("Helvetica", "B", 9)
        pdf.cell(45, 6, _safe_text(label + ":"), ln=0)
        pdf.set_font("Helvetica", "", 9)
        pdf.cell(pdf.epw - 45, 6, _safe_text(exec_b.get(key, "")), ln=True)
    pdf.ln(3)
    _write_section(pdf, "Síntesis", exec_b.get("resumen"))
    _footer(pdf, 2)

    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(pdf.epw, 9, "Resumen del escaneo", ln=True)
    pdf.ln(2)
    for label, key in (
        ("Equipos detectados", "equipos_detectados"),
        ("Equipos nuevos", "equipos_nuevos"),
        ("Equipos desconocidos", "equipos_desconocidos"),
        ("Endpoints", "endpoints"),
        ("Servicios analizados", "servicios_analizados"),
        ("Tiempo del analisis", "tiempo_analisis"),
    ):
        _write_section(pdf, label, str(scan.get(key, "")))
    _footer(pdf, 3)

    devices = pres.get("devices") or []
    if devices:
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(pdf.epw, 9, "Dispositivos analizados", ln=True)
        pdf.ln(2)
        cols = ["Nombre", "IP", "MAC", "Tipo", "Estado", "SO"]
        widths = [32, 28, 32, 22, 28, 38]
        _write_table_header(pdf, cols, widths)
        for d in devices[:25]:
            _write_table_row(pdf, [
                d.get("nombre"), d.get("ip"), d.get("mac"), d.get("tipo"),
                d.get("estado"), d.get("sistema_operativo"),
            ], widths)
        _footer(pdf, pdf.page_no())

    ports = pres.get("open_ports") or []
    if ports:
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(pdf.epw, 9, "Puertos abiertos", ln=True)
        pdf.ln(2)
        cols = ["Puerto", "Servicio", "Estado", "Riesgo", "Descripcion"]
        widths = [18, 35, 22, 22, 73]
        _write_table_header(pdf, cols, widths)
        for p in ports[:40]:
            _write_table_row(pdf, [
                p.get("puerto"), p.get("servicio"), p.get("estado"),
                p.get("nivel_riesgo"), p.get("descripcion"),
            ], widths)
        _footer(pdf, pdf.page_no())

    vulns = pres.get("vulnerabilities") or []
    if vulns:
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(pdf.epw, 9, "Vulnerabilidades", ln=True)
        for v in vulns[:20]:
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 10)
            pdf.cell(pdf.epw, 7, _safe_text(v.get("nombre")), ln=True)
            pdf.set_font("Helvetica", "", 9)
            _write_section(pdf, "Descripcion", v.get("descripcion"))
            _write_section(pdf, "Criticidad / Equipo", f"{v.get('criticidad')} — {v.get('equipo')}")
            _write_section(pdf, "Evidencias", v.get("evidencias"))
            _write_section(pdf, "Recomendacion", v.get("recomendacion"))
            if pdf.get_y() > 240:
                _footer(pdf, pdf.page_no())
                pdf.add_page()
        _footer(pdf, pdf.page_no())

    kernel = pres.get("kernel_analysis") or {}
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(pdf.epw, 9, "Analisis del Kernel IA", ln=True)
    _write_section(pdf, "Que encontro", kernel.get("que_encontro"))
    _write_section(pdf, "Que significa", kernel.get("que_significa"))
    _write_list(pdf, "Riesgos", kernel.get("riesgos") or [])
    _write_section(pdf, "Prioridad", kernel.get("prioridad"))
    _write_list(pdf, "Acciones recomendadas", kernel.get("acciones_recomendadas") or [])
    _footer(pdf, pdf.page_no())

    recs = pres.get("recommendations") or {}
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(pdf.epw, 9, "Recomendaciones", ln=True)
    for title, key in (("Alta prioridad", "alta"), ("Prioridad media", "media"), ("Prioridad baja", "baja")):
        items = recs.get(key) or []
        if items:
            _write_list(pdf, title, items)
    _write_list(pdf, "Evidencias", pres.get("evidences") or [])
    _footer(pdf, pdf.page_no())

    out_path = os.path.join(REPORTS_DIR, f"{report['id']}.pdf")
    pdf.safe_output(out_path)
    return out_path


def _generate_presentation_html(pres: dict, report: dict) -> str:
    out_path = os.path.join(REPORTS_DIR, f"{report['id']}.html")
    exec_b = pres.get("executive") or {}
    esc = lambda x: _safe_text(x)
    html = f"""<!DOCTYPE html><html lang="es"><head><meta charset="utf-8">
    <title>{esc(report.get('id'))} — NOVUS</title>
    <style>
    body{{font-family:'Segoe UI',Arial,sans-serif;margin:0;background:#f1f5f9;color:#0f172a}}
    .hero{{background:linear-gradient(135deg,#5d5fef,#4338ca);color:#fff;padding:2rem;text-align:center}}
    .wrap{{max-width:920px;margin:0 auto;padding:2rem}}
    h2{{color:#4338ca;border-bottom:2px solid #e2e8f0;padding-bottom:.5rem}}
    table{{width:100%;border-collapse:collapse;font-size:13px;margin:1rem 0}}
    th,td{{border:1px solid #cbd5e1;padding:6px 8px;text-align:left}}
    th{{background:#e2e8f0}}
    .card{{background:#fff;border-radius:8px;padding:1rem;margin:1rem 0;border:1px solid #e2e8f0}}
    </style></head><body>
    <div class="hero"><h1>NOVUS</h1><p>{esc(pres.get('title'))}</p></div>
    <div class="wrap">
    <h2>Resumen ejecutivo</h2>
    <div class="card"><p>{esc(exec_b.get('resumen'))}</p></div>
    </div></body></html>"""
    with open(out_path, "w", encoding="utf-8") as handle:
        handle.write(html)
    return out_path


def _generate_pdf_impl(report):
    return generate_presentation_pdf(report)


def _generate_pdf_impl_legacy(report):
    if report.get("report_type") == "enterprise":
        return generate_enterprise_pdf(report)
    if report.get("tipo") == "Centro de Defensa":
        return _generate_manual_defense_pdf(report)

    try:
        from fpdf import FPDF
    except ImportError:
        return _generate_html_fallback(report)

    from utils.novus_fpdf import create_novus_pdf

    pdf = create_novus_pdf()
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(pdf.epw, 10, "NOVUS Security Platform", ln=True)
    pdf.set_font("Helvetica", "I", 10)
    pdf.cell(pdf.epw, 6, _safe_text(NOVUS_SLOGAN), ln=True)
    pdf.ln(4)

    pdf.set_font("Helvetica", "B", 14)
    title = "Informe de Playbook" if report.get("tipo") == "Playbook" else "Informe de Vulnerabilidad"
    pdf.cell(pdf.epw, 9, _safe_text(title), ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(pdf.epw, 7, _safe_text(f"ID: {report.get('id', '')}"), ln=True)
    pdf.cell(pdf.epw, 7, _safe_text(f"Fecha: {report.get('fecha', '')}"), ln=True)
    pdf.cell(
        pdf.epw, 7,
        _safe_text(
            f"Severidad: {report.get('severidad', '')} | "
            f"Estado: {report.get('estado', '')} | "
            f"Confianza: {(report.get('technical') or {}).get('nivel_confianza', '')}"
        ),
        ln=True,
    )
    pdf.ln(4)

    _write_section(pdf, "Resumen ejecutivo", report.get("resumen_ejecutivo", ""))

    tech = report.get("technical", {})
    _write_section(pdf, "Descripcion", tech.get("descripcion_sencilla") or tech.get("que_ocurrio"))
    _write_section(pdf, "Descripcion tecnica", tech.get("descripcion_tecnica"))
    _write_section(pdf, "Evidencia", "\n".join(tech.get("evidencias") or []))
    _write_section(pdf, "Causa probable", tech.get("causa_probable") or tech.get("origen"))
    _write_section(pdf, "Impacto potencial", tech.get("impacto_potencial") or tech.get("consecuencias"))
    _write_section(pdf, "Nivel de riesgo", tech.get("riesgo_sistema"))
    _write_section(pdf, "Motor de deteccion", tech.get("motor"))
    _write_section(pdf, "Fuente de evidencia", tech.get("fuente_evidencia"))

    _write_list(pdf, "Acciones realizadas por NOVUS", tech.get("acciones_automaticas") or [])
    if tech.get("modificaciones_realizadas"):
        _write_list(pdf, "Modificaciones realizadas", tech.get("modificaciones_realizadas") or [])
    if tech.get("acciones_no_realizadas"):
        _write_list(pdf, "Acciones no realizadas o pendientes", tech.get("acciones_no_realizadas") or [])
    if tech.get("motores_utilizados"):
        _write_list(pdf, "Motores utilizados", tech.get("motores_utilizados") or [])
    _write_list(pdf, "Recomendaciones al cliente", tech.get("recomendaciones") or [])
    _write_list(pdf, "Acciones opcionales", tech.get("recomendaciones_opcionales") or [])

    ver = report.get("verification") or tech.get("verificacion") or {}
    if ver:
        _write_section(
            pdf, "Resultado de verificacion",
            f"Estado: {report.get('estado')} — {ver.get('detalle') or ver.get('detail')} "
            f"(Motor: {ver.get('motor')})"
        )

    rem_log = report.get("remediation_log") or []
    if rem_log:
        lines = [f"[{s.get('timestamp')}] {s.get('label')} — {s.get('detail') or ''}" for s in rem_log]
        _write_section(pdf, "Log de remediacion", "\n".join(lines))

    timeline = tech.get("linea_tiempo") or []
    if timeline:
        tl_text = "\n".join(
            f"{t.get('hora') or t.get('time')}: {t.get('evento') or t.get('event')}"
            for t in timeline[:15]
        )
        _write_section(pdf, "Historial", tl_text)

    _write_section(pdf, "Estado final", report.get("estado", ""))
    _write_section(pdf, "Tiempo de resolucion", report.get("tiempo_resolucion") or "Pendiente")
    _write_section(pdf, "Conclusiones", report.get("conclusiones", ""))

    pdf.ln(6)
    pdf.set_font("Helvetica", "I", 9)
    pdf.cell(pdf.epw, 6, "Informe generado automaticamente por NOVUS Security Platform.", ln=True)

    out_path = os.path.join(REPORTS_DIR, f"{report['id']}.pdf")
    pdf.safe_output(out_path)
    return out_path


def generate_pdf(report):
    """Return path to generated PDF file (con fallback HTML si falla fpdf)."""
    try:
        return _generate_pdf_impl(report)
    except Exception as exc:
        logger.error("generate_pdf falló para %s: %s", report.get("id"), exc, exc_info=True)
        return _generate_html_fallback(report)


def _generate_enterprise_html(report: dict) -> str:
    out_path = os.path.join(REPORTS_DIR, f"{report['id']}.html")
    cover = report.get("cover") or {}
    exec_sum = report.get("executive_summary") or {}
    sig = report.get("signature") or {}
    charts = report.get("charts") or {}
    details = report.get("technical_details") or []

    chart_html = ""
    for title, data in charts.items():
        if isinstance(data, dict):
            bars = "".join(
                f'<div class="bar-row"><span>{k}</span><div class="bar" style="width:{min(100, (v/max(data.values())*100) if max(data.values()) else 0)}%"></div><span>{v}</span></div>'
                for k, v in data.items() if isinstance(v, (int, float))
            )
            chart_html += f'<h3>{title}</h3>{bars}'

    findings_html = "".join(
        f'<div class="finding"><h4>{d.get("id")}</h4><p>{d.get("descripcion","")}</p>'
        f'<p><b>Motor:</b> {d.get("motor")} | <b>Riesgo:</b> {d.get("riesgo")} | <b>Estado:</b> {d.get("estado")}</p></div>'
        for d in details[:20]
    )

    html = f"""<!DOCTYPE html><html lang="es"><head><meta charset="utf-8">
    <title>{report.get('id')} — NOVUS</title>
    <style>
    body{{font-family:'Segoe UI',Arial,sans-serif;margin:0;color:#111;background:#f8fafc}}
    .cover{{background:linear-gradient(135deg,#5d5fef,#4338ca);color:#fff;padding:60px 40px;text-align:center}}
    .cover h1{{font-size:42px;margin:0}} .cover h2{{font-weight:400;opacity:.9}}
    .content{{max-width:900px;margin:0 auto;padding:40px 24px}}
    h2{{color:#5d5fef;border-bottom:2px solid #e2e8f0;padding-bottom:8px}}
    .meta{{display:grid;grid-template-columns:1fr 1fr;gap:8px;font-size:14px;margin:20px 0}}
    .bar-row{{display:flex;align-items:center;gap:8px;margin:4px 0;font-size:12px}}
    .bar{{height:14px;background:#5d5fef;border-radius:3px;min-width:2px}}
    .finding{{background:#fff;border:1px solid #e2e8f0;border-radius:8px;padding:12px;margin:8px 0}}
    .footer{{text-align:center;padding:24px;color:#64748b;font-size:12px;border-top:1px solid #e2e8f0;margin-top:40px}}
    </style></head><body>
    <div class="cover">
        <h1>NOVUS</h1><h2>{cover.get('titulo','Informe de Seguridad')}</h2>
        <p>{cover.get('empresa')} · {cover.get('sector')}</p>
        <p>{cover.get('fecha')} {cover.get('hora')} · {report.get('id')} · v{report.get('version')}</p>
        <p><strong>Nivel de seguridad: {cover.get('nivel_seguridad')}</strong></p>
    </div>
    <div class="content">
        <h2>Resumen Ejecutivo</h2>
        <p>{exec_sum.get('texto_ejecutivo','')}</p>
        <div class="meta">
            <div>Estado: <b>{exec_sum.get('estado_general')}</b></div>
            <div>Vulnerabilidades: <b>{exec_sum.get('vulnerabilidades')}</b></div>
            <div>Amenazas: <b>{exec_sum.get('amenazas')}</b></div>
            <div>Dispositivos: <b>{exec_sum.get('dispositivos_monitoreados')}</b></div>
        </div>
        <h2>Indicadores</h2>{chart_html}
        <h2>Detalle Tecnico</h2>{findings_html or '<p>Sin hallazgos activos.</p>'}
    </div>
    <div class="footer">
        <p>{NOVUS_SLOGAN}</p>
        <p>Generado por {sig.get('generado_por')} · {sig.get('fecha')} {sig.get('hora')} · v{sig.get('version')}</p>
        <p>SHA-256: {sig.get('integridad_sha256','')}</p>
    </div></body></html>"""
    with open(out_path, "w", encoding="utf-8") as handle:
        handle.write(html)
    return out_path


def _generate_html_fallback(report):
    if report.get("tipo") == "Centro de Defensa":
        from services.manual_defense_report_html import render_manual_defense_report_html

        out_path = os.path.join(REPORTS_DIR, f"{report['id']}.html")
        with open(out_path, "w", encoding="utf-8") as handle:
            handle.write(render_manual_defense_report_html(report))
        return out_path
    if report.get("report_type") == "enterprise":
        return _generate_enterprise_html(report)
    out_path = os.path.join(REPORTS_DIR, f"{report['id']}.html")
    tech = report.get("technical", {})
    recs = "".join(f"<li>{r}</li>" for r in tech.get("recomendaciones", []))
    html = f"""<!DOCTYPE html><html><head><meta charset='utf-8'><title>{report.get('id')}</title>
    <style>body{{font-family:Arial,sans-serif;margin:40px;color:#111}}h1{{color:#5d5fef}}
    .meta{{color:#555;font-size:14px}}h2{{margin-top:24px;font-size:16px}}</style></head><body>
    <h1>NOVUS Security Platform</h1><p><em>{NOVUS_SLOGAN}</em></p>
    <h2>Informe de Vulnerabilidad</h2>
    <p class='meta'><b>ID:</b> {report.get('id')}<br><b>Fecha:</b> {report.get('fecha')}<br>
    <b>Severidad:</b> {report.get('severidad')} | <b>Estado:</b> {report.get('estado')}</p>
    <h2>Resumen ejecutivo</h2><p>{report.get('resumen_ejecutivo','')}</p>
    <h2>Causa probable</h2><p>{tech.get('causa_probable') or tech.get('origen','')}</p>
    <h2>Impacto</h2><p>{tech.get('impacto_potencial') or tech.get('consecuencias','')}</p>
    <h2>Recomendaciones</h2><ul>{recs}</ul>
    <p><em>Informe generado automaticamente por NOVUS Security Platform.</em></p>
    </body></html>"""
    with open(out_path, "w", encoding="utf-8") as handle:
        handle.write(html)
    logger.warning("fpdf2 no instalado — informe exportado como HTML")
    return out_path


def export_report_json(report: dict) -> str:
    out_path = os.path.join(REPORTS_DIR, f"{report['id']}.export.json")
    with open(out_path, "w", encoding="utf-8") as handle:
        json.dump(report, handle, indent=2, ensure_ascii=False)
    return out_path


def export_report_csv(report: dict) -> str:
    out_path = os.path.join(REPORTS_DIR, f"{report['id']}.csv")
    rows = report.get("technical_details") or []
    if not rows:
        tech = report.get("technical") or {}
        rows = [tech] if tech else [{"id": report.get("id"), "descripcion": report.get("resumen_ejecutivo")}]

    fieldnames = ["id", "descripcion", "motor", "confianza", "riesgo", "estado", "impacto", "fecha", "hora"]
    with open(out_path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in fieldnames})
    return out_path


def export_report_xlsx(report: dict) -> str:
    """Excel (.xlsx) vía openpyxl; si no está instalado, CSV con BOM para Excel."""
    out_path = os.path.join(REPORTS_DIR, f"{report['id']}.xlsx")
    rows = report.get("technical_details") or []
    if not rows:
        tech = report.get("technical") or {}
        rows = [tech] if tech else [{"id": report.get("id"), "descripcion": report.get("resumen_ejecutivo")}]
    fieldnames = ["id", "descripcion", "motor", "confianza", "riesgo", "estado", "impacto", "fecha", "hora"]
    try:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Informe"
        ws.append(fieldnames)
        for row in rows:
            ws.append([row.get(k, "") for k in fieldnames])
        meta = report.get("metadata") or {}
        ws.append([])
        ws.append(["Campo", "Valor"])
        for k, v in meta.items():
            ws.append([k, str(v)[:500]])
        wb.save(out_path)
        return out_path
    except ImportError:
        csv_path = export_report_csv(report)
        xlsx_compat = os.path.join(REPORTS_DIR, f"{report['id']}.xlsx")
        with open(csv_path, "rb") as src, open(xlsx_compat, "wb") as dst:
            dst.write(b"\xef\xbb\xbf")
            dst.write(src.read())
        return xlsx_compat


def _write_forensic_export_manifest(report: dict, path: str) -> None:
    try:
        from services.forensic_evidence_integrity_service import build_export_integrity_manifest

        manifest = build_export_integrity_manifest(report, path)
        manifest_path = f"{path}.forensic_manifest.json"
        with open(manifest_path, "w", encoding="utf-8") as mf:
            json.dump(manifest, mf, indent=2, ensure_ascii=False)
    except Exception as exc:
        logger.debug("forensic export manifest: %s", exc)


def export_report(report: dict, fmt: str = "pdf") -> str:
    fmt = (fmt or "pdf").lower()
    path: str
    if fmt == "pdf":
        try:
            path = generate_pdf(report)
        except Exception as exc:
            logger.error("PDF export failed for %s: %s", report.get("id"), exc, exc_info=True)
            if report.get("tipo") == "Centro de Defensa":
                from services.manual_defense_report_html import render_manual_defense_report_html

                out_path = os.path.join(REPORTS_DIR, f"{report['id']}.html")
                with open(out_path, "w", encoding="utf-8") as handle:
                    handle.write(render_manual_defense_report_html(report))
                path = out_path
            else:
                path = _generate_html_fallback(report)
    elif fmt == "html":
        path = _generate_enterprise_html(report) if report.get("report_type") == "enterprise" else _generate_html_fallback(report)
    elif fmt == "json":
        path = export_report_json(report)
    elif fmt == "csv":
        path = export_report_csv(report)
    elif fmt in ("xlsx", "excel"):
        path = export_report_xlsx(report)
    else:
        raise ValueError(f"Formato no soportado: {fmt}")
    _write_forensic_export_manifest(report, path)
    return path


INTEL_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "intelligence")


def generate_case_pdf(case: dict) -> str:
    """Exporta expediente del Threat Intelligence Center a PDF."""
    os.makedirs(INTEL_DIR, exist_ok=True)
    try:
        from fpdf import FPDF
    except ImportError:
        out_path = os.path.join(INTEL_DIR, f"{case['id']}.html")
        with open(out_path, "w", encoding="utf-8") as handle:
            handle.write(f"<html><body><h1>Caso {case.get('id')}</h1><pre>{case}</pre></body></html>")
        return out_path

    pdf = FPDF()
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(pdf.epw, 10, "NOVUS Threat Intelligence Center", ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(pdf.epw, 7, _safe_text(f"Caso: {case.get('id', '')}"), ln=True)
    pdf.cell(pdf.epw, 7, _safe_text(f"Fecha: {case.get('fecha', '')} {case.get('hora', '')}"), ln=True)
    pdf.cell(pdf.epw, 7, _safe_text(
        f"Tipo: {case.get('tipo', '')} | Riesgo: {case.get('nivel_riesgo', '')} | Estado: {case.get('estado', '')}"
    ), ln=True)
    pdf.ln(4)

    lessons = case.get("lecciones_aprendidas") or {}
    if isinstance(lessons, dict):
        _write_section(pdf, "Resumen ejecutivo", lessons.get("que_ocurrio", ""))
        _write_section(pdf, "Como ocurrio", lessons.get("como_ocurrio", ""))
        _write_section(pdf, "Respuesta NOVUS", lessons.get("respuesta_novus", ""))

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(pdf.epw, 8, "Cronologia", ln=True)
    pdf.set_font("Helvetica", "", 9)
    for step in case.get("timeline") or []:
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(pdf.epw, 5, _safe_text(
            f"{step.get('timestamp', '')} — {step.get('evento', '')}: {step.get('detalle', '')}"
        ))

    _write_section(pdf, "Conclusiones", case.get("resultado", ""))
    out_path = os.path.join(INTEL_DIR, f"{case['id']}.pdf")
    pdf.output(out_path)
    return out_path
